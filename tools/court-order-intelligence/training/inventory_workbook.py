"""Read-only private acquisition inventory. NEVER a supervised/gold-data exporter.

Office notes, titles, advocates, addresses, awards and directions are deliberately
not copied. Even this minimized register inventory stays local-private: official
PDFs must independently establish every eventual public training proposition.
"""
import argparse
from collections import Counter
import hashlib
import json
from pathlib import Path
import re

COURTS = {"delhi high court": "DelhiHighCourt", "high court": "CourtNeedsConfirmation"}
STATUSES = {"pending": "Pending", "disposed off": "Disposed", "diposed off": "Disposed"}
REQUIRED = {"case satus", "case no.", "which court pertains to", "last order link"}


def clean(value):
    return " ".join(str(value).split()) if value is not None else ""


def sha256(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def select_rows(sheet, values):
    """Use explicit court/status aliases; never infer a court from case spelling."""
    columns = None
    selected, exceptions = [], []
    for row_number, row in enumerate(values, 1):
        names = [clean(x).casefold() for x in row]
        if REQUIRED.issubset(names):
            if columns is not None:
                raise ValueError("Repeated register header requires manual review")
            columns = {name: names.index(name) for name in REQUIRED}
            continue
        if columns is None:
            continue

        def cell(name):
            i = columns[name]
            return clean(row[i]) if i < len(row) else ""

        court_raw, status_raw = cell("which court pertains to"), cell("case satus")
        court = COURTS.get(court_raw.casefold())
        if court is None:
            continue
        status = STATUSES.get(status_raw.casefold())
        if status is None:
            exceptions.append({"sheet": sheet, "excel_row": row_number,
                               "raw_court": court_raw, "raw_status": status_raw,
                               "reason": "MissingOrUnsupportedStatus"})
            continue
        case = cell("case no.")
        # Identity handling is deliberately conservative: raw register spelling
        # remains intact; no first-number guess for multi-case cells.
        issues = []
        if court == "CourtNeedsConfirmation":
            issues.append("CourtNeedsConfirmation")
        if not case:
            issues.append("MissingCaseNumber")
        elif re.search(r"_{2,}|\.{3,}|/\s*\d{4}$", case) and not re.search(r"\d+\s*/\s*\d{4}", case):
            issues.append("IncompleteCaseNumber")
        key = re.sub(r"\s+", " ", case.casefold()).strip()
        selected.append({"sheet": sheet, "excel_row": row_number,
                         "raw_court": court_raw, "court_state": court,
                         "raw_status": status_raw, "register_status": status,
                         "raw_case_number": case, "candidate_order_url": cell("last order link"),
                         "annotation_state": "UNREVIEWED", "issues": issues,
                         "raw_duplicate_key": key})
    if columns is None:
        raise ValueError("Court register header not found")
    duplicates = Counter(x["raw_duplicate_key"] for x in selected if x["raw_duplicate_key"])
    for item in selected:
        if duplicates[item["raw_duplicate_key"]] > 1:
            item["issues"].append("RepeatedRawCaseNumber")
    return selected, exceptions


def inventory(path):
    import openpyxl
    before = sha256(path)
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows, exceptions = [], []
    try:
        for sheet in workbook:
            # Inspect headers only on unrelated sheets; never export their data.
            headers = list(sheet.iter_rows(min_row=1, max_row=5, values_only=True))
            if not any(REQUIRED.issubset({clean(x).casefold() for x in row}) for row in headers):
                continue
            selected, excluded = select_rows(sheet.title, sheet.iter_rows(values_only=True))
            rows.extend(selected)
            exceptions.extend(excluded)
    finally:
        workbook.close()
    if sha256(path) != before:
        raise ValueError("Workbook changed during extraction; inventory discarded")
    counts = Counter((r["court_state"], r["register_status"]) for r in rows)
    return {"schema_version": 1, "private_acquisition_inventory": True,
            "training_eligible": False, "source_filename": Path(path).name,
            "workbook_sha256": before, "rows": rows, "exceptions": exceptions,
            "summary": {"selected_rows": len(rows), "exceptions": len(exceptions),
                        "counts": [{"court_state": c, "register_status": s, "rows": n}
                                   for (c, s), n in sorted(counts.items())]}}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    args = parser.parse_args()
    result = inventory(args.workbook)
    target = Path(__file__).parent / "local-private" / "workbook-inventory.json"
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"path": str(target), **result["summary"]}))


if __name__ == "__main__":
    main()
